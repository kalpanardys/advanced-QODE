from docx import Document
from pptx import Presentation

def export_to_word(query, result, filename="impact_report.docx"):
    doc = Document()

    doc.add_heading("Agentic AI Process Intelligence Report", level=1)

    doc.add_heading("Query", level=2)
    doc.add_paragraph(query)

    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(result.get("summary", ""))

    doc.add_heading("Impacted Entities", level=2)
    for entity in result.get("impacted_entities", []):
        doc.add_paragraph(entity, style="List Bullet")

    doc.add_heading("Risks", level=2)
    for risk in result.get("risks", []):
        doc.add_paragraph(risk, style="List Bullet")

    doc.add_heading("Recommendations", level=2)
    for rec in result.get("recommendations", []):
        doc.add_paragraph(rec, style="List Bullet")

    doc.save(filename)

    print(f"\n[EXPORT] Word report generated: {filename}")

def export_to_ppt(query, result, filename="impact_report.pptx"):
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Agentic AI Impact Analysis"

    slide.placeholders[1].text = query

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"

    slide.placeholders[1].text = result.get("summary", "")

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Impacted Entities"

    content = "\n".join(result.get("impacted_entities", []))
    slide.placeholders[1].text = content

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recommendations"

    content = "\n".join(result.get("recommendations", []))
    slide.placeholders[1].text = content

    prs.save(filename)

    print(f"\n[EXPORT] PowerPoint report generated: {filename}")