from docx import Document
from pptx import Presentation


class ReportExporter:

    @staticmethod
    def export_word(
        query,
        entity,
        graph_evidence,
        semantic_evidence,
        llm_response,
        filename="DQODE_Analysis.docx"
    ):

        doc = Document()

        doc.add_heading(
            "DQODE Graph RAG Analysis Report",
            level=1
        )

        doc.add_heading("Question", level=2)
        doc.add_paragraph(query)

        doc.add_heading("Detected Entity", level=2)
        doc.add_paragraph(entity)

        doc.add_heading("Graph Evidence", level=2)
        doc.add_paragraph(graph_evidence)

        doc.add_heading("Semantic Evidence", level=2)

        for item in semantic_evidence:
            doc.add_paragraph(item)

        doc.add_heading("LLM Analysis", level=2)
        doc.add_paragraph(llm_response)

        doc.save(filename)

        print(f"Word report created: {filename}")

    @staticmethod
    def export_ppt(
        query,
        entity,
        graph_evidence,
        semantic_evidence,
        llm_response,
        filename="DQODE_Analysis.pptx"
    ):

        prs = Presentation()

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = "DQODE Graph RAG Analysis"

        slide.placeholders[1].text = (
            f"Question:\n{query}\n\n"
            f"Entity:\n{entity}"
        )

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = "Graph Evidence"

        slide.placeholders[1].text = graph_evidence

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = "Semantic Evidence"

        slide.placeholders[1].text = "\n".join(
            semantic_evidence
        )

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = "LLM Analysis"

        slide.placeholders[1].text = llm_response[:3000]

        prs.save(filename)

        print(f"PPT report created: {filename}")